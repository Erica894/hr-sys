"""BudgetOverride 批量导入 service: 解析 Excel + 校验 + 写库。

Excel 列 (固定顺序, 含表头):
  department_code, department_name, adjustment_type(ANNUAL/PROMOTION), override_amount_cny

校验规则:
  - department_code 必须存在且 OrgUnit.type=DEPT
  - adjustment_type ∈ {ANNUAL, PROMOTION}
  - override_amount_cny >= 0
  - override 不能低于该 (cycle, dept, adj_type) 已分配额 (来自 AdjustmentProposal)
任一行失败 → 整批拒绝，回报行级错误清单。
"""
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook

from apps.compensation_plan.models import BudgetOverride
from apps.compensation_plan.services.budget_aggregation import (
    aggregate_adjustment_allocated,
)
from apps.iam.models import OrgUnit


COLUMNS = [
    ("department_code", "部门 code"),
    ("department_name", "部门名称（仅展示）"),
    ("adjustment_type", "调薪类型 (ANNUAL/PROMOTION)"),
    ("override_amount_cny", "覆盖金额 CNY (月度)"),
]
VALID_ADJ = {"ANNUAL", "PROMOTION"}


class ImportError(Exception):
    """批量导入校验错误，payload=行级错误清单。"""

    def __init__(self, errors: list[dict]):
        self.errors = errors
        super().__init__(f"{len(errors)} row(s) failed")


def build_template_xlsx() -> bytes:
    """生成空白模板 (含表头 + 1 行示例)，返回 xlsx bytes。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "BudgetOverride"
    ws.append([label for _, label in COLUMNS])
    ws.append(["ENG", "Engineering", "ANNUAL", 6000.00])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _allocated_by_dept_adj(cycle) -> dict[tuple[int, str], Decimal]:
    """按 (dept_id, adj_type) 聚合已分配额 (跨 cat1 累加)。"""
    raw = aggregate_adjustment_allocated(cycle)
    out: dict[tuple[int, str], Decimal] = defaultdict(lambda: Decimal("0"))
    for (ou_id, adj, _cat), amount in raw.items():
        out[(ou_id, adj)] += amount
    return dict(out)


def _parse_rows(file_bytes: bytes) -> list[dict]:
    """读取 xlsx，返回 [{row_index, raw}, ...]。表头在第 1 行，数据从第 2 行起。"""
    wb = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb.active
    rows = []
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if idx == 1:
            continue
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        padded = list(row) + [None] * (len(COLUMNS) - len(row))
        rows.append({
            "row_index": idx,
            "department_code": (str(padded[0]).strip() if padded[0] is not None else ""),
            "department_name": (str(padded[1]).strip() if padded[1] is not None else ""),
            "adjustment_type": (str(padded[2]).strip().upper() if padded[2] is not None else ""),
            "override_amount_cny": padded[3],
        })
    return rows


def import_overrides(cycle, file_bytes: bytes, user) -> dict:
    """解析 + 校验 + 写库。失败抛 ImportError。

    返回 {created: int, updated: int, total_rows: int}。
    """
    parsed = _parse_rows(file_bytes)
    errors: list[dict] = []

    dept_by_code = {
        u.code: u for u in OrgUnit.objects.filter(type="DEPT")
    }
    allocated = _allocated_by_dept_adj(cycle)

    valid_payloads: list[tuple[OrgUnit, str, Decimal]] = []
    for r in parsed:
        idx = r["row_index"]
        code = r["department_code"]
        adj = r["adjustment_type"]
        amt_raw = r["override_amount_cny"]

        if not code:
            errors.append({"row": idx, "field": "department_code", "msg": "必填"})
            continue
        dept = dept_by_code.get(code)
        if dept is None:
            errors.append({"row": idx, "field": "department_code", "msg": f"找不到 type=DEPT 的部门 code={code}"})
            continue
        if adj not in VALID_ADJ:
            errors.append({"row": idx, "field": "adjustment_type", "msg": f"必须是 ANNUAL 或 PROMOTION, got {adj!r}"})
            continue
        try:
            amount = Decimal(str(amt_raw)) if amt_raw is not None else None
        except (InvalidOperation, ValueError):
            errors.append({"row": idx, "field": "override_amount_cny", "msg": f"金额格式不合法 {amt_raw!r}"})
            continue
        if amount is None or amount < 0:
            errors.append({"row": idx, "field": "override_amount_cny", "msg": "金额必须 ≥ 0"})
            continue
        already = allocated.get((dept.id, adj), Decimal("0"))
        if amount < already:
            errors.append({
                "row": idx, "field": "override_amount_cny",
                "msg": f"覆盖金额 {amount} 低于已分配额 {already.quantize(Decimal('0.01'))}",
            })
            continue
        valid_payloads.append((dept, adj, amount.quantize(Decimal("0.01"))))

    if errors:
        raise ImportError(errors)

    created = 0
    updated = 0
    with transaction.atomic():
        for dept, adj, amount in valid_payloads:
            obj, was_created = BudgetOverride.objects.update_or_create(
                reward_cycle=cycle, department=dept, adjustment_type=adj,
                defaults={
                    "override_amount_cny": amount,
                    "source": "IMPORTED",
                    "uploaded_by": user if user and user.is_authenticated else None,
                },
            )
            if was_created:
                created += 1
            else:
                updated += 1

    return {"created": created, "updated": updated, "total_rows": len(valid_payloads)}
